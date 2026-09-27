"""Render the full Morpheus Pharmacogenomics Report as an editable Word document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .fonts import ASSETS
from .phenotypes import GREEN, RED, YELLOW
from .report_text import (HOW_TO_READ, QUICK_REF_NOTE, REPORT_GUIDE, SHARE_NOTE, footer_text,
                          interpretation_notes, patient_rows)
from .rules import GENE_KEY_DRUGS
from .summary import Summary

NAVY = RGBColor(0x0A, 0x16, 0x28)
FOOT = RGBColor(0x7F, 0x7F, 0x7F)
TIER_HEX = {RED: "FF0000", YELLOW: "FFC000", GREEN: "00B050"}
FONT = "Calibri"


# --------------------------------------------------------------------------- low-level helpers

def _run(p, text, bold=False, italic=False, size=None, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def _rich(p, markup: str, size=None):
    """Add text containing <b>..</b> spans."""
    for i, part in enumerate(re.split(r"</?b>", markup)):
        if part:
            _run(p, part, bold=(i % 2 == 1), size=size)


def _para(doc_or_cell, text="", bold=False, italic=False, size=11, color=None, align=None,
          space_after=4, space_before=0):
    p = doc_or_cell.add_paragraph()
    if text:
        _run(p, text, bold, italic, size, color)
    pf = p.paragraph_format
    pf.space_after, pf.space_before = Pt(space_after), Pt(space_before)
    if align:
        p.alignment = align
    return p


def _cell_text(cell, text, bold=False, size=8.5, align=None):
    p = cell.paragraphs[0]
    _run(p, text, bold=bold, size=size)
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _borders(cell, color="808080", sz=6, **sides):
    """Set cell borders; sides override e.g. left=("FF0000", 24)."""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        col, size = sides.get(side, (color, sz))
        el = b.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            b.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), col)


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _margins(cell, top=60, left=90, bottom=60, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        m.append(el)
    tcPr.append(m)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    trPr.append(el)


def _table(doc, rows, widths):
    t = doc.add_table(rows=rows, cols=len(widths))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for row in t.rows:
        _no_split(row)
        for cell, w in zip(row.cells, widths):
            cell.width = w
            _margins(cell)
    return t


def _heading(doc, text, rule=True, size=13, space_before=6):
    p = _para(doc, text, bold=True, size=size, color=NAVY, space_before=space_before, space_after=6)
    if rule:
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "18"), ("space", "2"), ("color", "0A1628")):
            bottom.set(qn(f"w:{k}"), v)
        bdr.append(bottom)
        pPr.append(bdr)
    p.paragraph_format.keep_with_next = True
    return p


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _header_cells(t, labels):
    row = t.rows[0]
    _repeat_header(row)
    for cell, label in zip(row.cells, labels):
        _cell_text(cell, label, bold=True)
        _shade(cell, "F2F2F2")
        _borders(cell, top=("0A1628", 18))


def _tier_row(row, tier):
    """Outline a row in its tier colour (thick outer edges, thin inner grid)."""
    color = TIER_HEX[tier]
    n = len(row.cells)
    for i, cell in enumerate(row.cells):
        sides = {"top": (color, 18), "bottom": (color, 18)}
        if i == 0:
            sides["left"] = (color, 18)
        if i == n - 1:
            sides["right"] = (color, 18)
        _borders(cell, **sides)


# Word rejects files whose property children are out of schema order, so sort them before saving.
_TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
               "textDirection", "tcFitText", "vAlign", "hideMark"]
_PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr",
              "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap",
              "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
              "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc",
              "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
              "sectPr", "pPrChange"]


def _sort_children(el, order):
    rank = {qn(f"w:{n}"): i for i, n in enumerate(order)}
    kids = list(el)
    kids.sort(key=lambda k: rank.get(k.tag, len(order)))
    for k in kids:
        el.remove(k)
    for k in kids:
        el.append(k)


def _normalize_xml(doc):
    body = doc.element.body
    for tcPr in body.iter(qn("w:tcPr")):
        _sort_children(tcPr, _TCPR_ORDER)
    for pPr in body.iter(qn("w:pPr")):
        _sort_children(pPr, _PPR_ORDER)


# --------------------------------------------------------------------------- sections

def _cover(doc, summary: Summary):
    p = _para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    p.add_run().add_picture(str(ASSETS / "logo_wide.png"), width=Inches(3.2))
    _para(doc, "PHARMACOGENOMICS REPORT", bold=True, size=16, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER,
          space_after=10)

    rows = patient_rows(summary)
    widths = [Inches(1.1), Inches(2.65), Inches(1.1), Inches(2.65)]
    t = _table(doc, len(rows), widths)
    for r, vals in zip(t.rows, rows):
        for i, (cell, v) in enumerate(zip(r.cells, vals)):
            _cell_text(cell, v, bold=(i % 2 == 0), size=10)
            _borders(cell)
    _para(doc, space_after=6)

    legend = [("Normal: standard dosing", GREEN), ("Intermediate: adjust or monitor", YELLOW),
              ("Rapid or poor: avoid or change", RED)]
    t = _table(doc, 1, [Inches(2.5)] * 3)
    for cell, (label, tier) in zip(t.rows[0].cells, legend):
        _cell_text(cell, label, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        _borders(cell, left=(TIER_HEX[tier], 36))
    _para(doc, space_after=8)

    _heading(doc, "How to Read This Report")
    for text in HOW_TO_READ:
        p = _para(doc, space_after=6)
        _rich(p, text, size=11)
    _para(doc, "Report guide", bold=True, size=12, color=NAVY, space_before=4, space_after=3)
    for label, desc in REPORT_GUIDE(summary):
        p = _para(doc, space_after=2)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        _run(p, "•  ", size=11)
        _run(p, f"{label}: ", bold=True, size=11)
        _run(p, desc, size=11)
    _para(doc, SHARE_NOTE, bold=True, size=11, space_before=8)
    _page_break(doc)


def _gene_results(doc, summary: Summary):
    _heading(doc, "Gene Results", rule=False)
    genes = summary.genes_sorted
    widths = [Inches(1.1), Inches(1.3), Inches(1.9), Inches(3.2)]
    t = _table(doc, len(genes) + 1, widths)
    _header_cells(t, ["Gene", "Genotype", "Phenotype", "Key drugs affected"])
    footnote = False
    for row, g in zip(t.rows[1:], genes):
        pheno = g.phenotype
        if g.name == "CYP3A5" and g.code == "PM":
            pheno += "*"
            footnote = True
        _cell_text(row.cells[0], g.name, bold=True)
        _cell_text(row.cells[1], g.genotype)
        _cell_text(row.cells[2], pheno, bold=True)
        _cell_text(row.cells[3], GENE_KEY_DRUGS.get(g.name, ""))
        _tier_row(row, g.tier)
    if footnote:
        _para(doc, "*CYP3A5 *3/*3 is the most common genotype in people of European ancestry and means "
                   "standard tacrolimus dosing.", italic=True, size=8, space_before=3)
    _page_break(doc)


def _quick_reference(doc, summary: Summary):
    _heading(doc, "Medication Quick Reference")
    heads = [("AVOID OR CHANGE", RED), ("ADJUST OR MONITOR", YELLOW), ("STANDARD DOSING", GREEN)]
    t = _table(doc, 2, [Inches(2.5)] * 3)
    for i, (label, tier) in enumerate(heads):
        hc, bc = t.rows[0].cells[i], t.rows[1].cells[i]
        _cell_text(hc, label, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        _borders(hc, left=(TIER_HEX[tier], 36))
        _borders(bc)
        bc.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        groups = summary.by_category((tier,))
        first = True
        for cat, rows in groups.items():
            p = bc.paragraphs[0] if first else bc.add_paragraph()
            first = False
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            _run(p, cat, bold=True, size=11, color=NAVY)
            p = bc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            _run(p, ", ".join(r.drug.name for r in rows), size=10.5)
        if first:
            _run(bc.paragraphs[0], "None", size=10.5)
    _para(doc, QUICK_REF_NOTE, italic=True, size=7.5, space_before=3)
    _page_break(doc)


def _drug_tables(doc, summary: Summary, tiers, title, subtitle=""):
    _para(doc, title, bold=True, size=14, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    if subtitle:
        _para(doc, subtitle, italic=True, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    widths = [Inches(1.3), Inches(1.3), Inches(3.85), Inches(1.05)]
    for cat, rows in summary.by_category(tiers).items():
        _heading(doc, cat, rule=False, size=12, space_before=8)
        t = _table(doc, len(rows) + 1, widths)
        _header_cells(t, ["Drug", "Gene(s)", "Recommendation", "Source"])
        for row, r in zip(t.rows[1:], rows):
            _cell_text(row.cells[0], r.drug.name, bold=True)
            _cell_text(row.cells[1], ", ".join(r.drug.genes))
            _cell_text(row.cells[2], r.rec.text, bold=True)
            _cell_text(row.cells[3], r.source)
            _tier_row(row, r.tier)


def _notes(doc, summary: Summary):
    _heading(doc, "Limitations & Interpretation Notes", space_before=14)
    _para(doc, "Interpretation notes", bold=True, size=8.5, space_after=3)
    for n in interpretation_notes(summary):
        p = _para(doc, space_after=3)
        _run(p, "•  " + n, size=8.5)


def _footer(doc, summary: Summary):
    sec = doc.sections[0]
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(p, footer_text(summary) + "  |  Page ", size=7.5, color=FOOT)
    r = _run(p, "", size=7.5, color=FOOT)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)


def render_full_docx(summary: Summary, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.5))
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    doc.core_properties.title = f"Pharmacogenomics Report – {summary.report.patient.display_name}"
    doc.core_properties.author = "Morpheus Precision Health"

    _cover(doc, summary)
    _gene_results(doc, summary)
    _quick_reference(doc, summary)
    if summary.by_category((RED, YELLOW)):
        _drug_tables(doc, summary, (RED, YELLOW), "ACTION REQUIRED: MEDICATIONS NEEDING ADJUSTMENT",
                     "Red and yellow medications only, grouped by system. Full list of all medications "
                     "follows.")
        _page_break(doc)
    _drug_tables(doc, summary, (RED, YELLOW, GREEN), "ALL MEDICATIONS: DETAILED RECOMMENDATIONS")
    _notes(doc, summary)
    _footer(doc, summary)
    _normalize_xml(doc)
    doc.save(str(out_path))
    return out_path
